"""Synthetic documents for tests and smoke runs (generated in a temporary directory; never committed).

``make_pdf`` builds a PDF with PyMuPDF: a text page, a vector-drawing page, a raster page (an image covering the
page), an empty page, a CP1251-mojibake page and a rotated page. ``make_epub``/``make_docx`` build minimal containers
with the structures the extractors look for. ``make_register`` writes a SOURCE_REGISTER.csv for a synthetic
resources root. All text is synthetic.
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
import shutil
import subprocess
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

TEXT_RU = ("Синтетический абзац для проверки извлечения текста и порядка чтения блоков. Он не взят из источников "
           "корпуса и нужен только тестам платформы.")
TEXT_EN = "Synthetic paragraph used only by the platform tests; it is not taken from any corpus source."
MOJIBAKE = "Ñèíòåòè÷åñêèé òåêñò â ñòàðîé êîäèðîâêå äëÿ ïðîâåðêè ðåìîíòà ñòðàíèöû"  # CP1251 read as Latin-1


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    h.update(Path(path).read_bytes())
    return h.hexdigest()


@lru_cache(maxsize=1)
def font_file() -> str | None:
    """A TrueType font with Cyrillic glyphs for synthetic pages, without machine paths in code.

    Order: ``VKM_FONT_FILE``; fontconfig (``fc-match`` for ``lang=ru``); the OS font directory named by the
    environment (``WINDIR``/``SystemRoot``). ``None`` when nothing is found (PyMuPDF's base font is used then)."""
    env = os.environ.get("VKM_FONT_FILE", "").strip()
    if env and Path(env).is_file():
        return env
    if shutil.which("fc-match"):
        for family in ("DejaVu Sans", "Liberation Sans", "Arial"):
            try:
                out = subprocess.run(["fc-match", "-f", "%{file}", f"{family}:lang=ru"], capture_output=True,
                                     text=True, timeout=20).stdout.strip()
            except (OSError, subprocess.SubprocessError):
                out = ""
            if out.lower().endswith((".ttf", ".otf")) and Path(out).is_file():
                return out
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windir:
        for name in ("arial.ttf", "times.ttf"):
            p = Path(windir) / "Fonts" / name
            if p.is_file():
                return str(p)
    return None


_font = font_file


def make_png(width: int = 400, height: int = 300, *, text: str | None = None, color: int = 90) -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("L", (width, height), 255)
    d = ImageDraw.Draw(img)
    d.rectangle([10, 10, width - 10, height - 10], outline=color, width=4)
    d.line([10, 10, width - 10, height - 10], fill=color, width=3)
    if text:
        d.text((20, height // 2), text, fill=0)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_pdf(path: Path, *, pages: tuple[str, ...] = ("text", "vector", "raster", "empty", "mojibake", "rotated"),
             bom: bool = False, labels: bool = False) -> Path:
    """PDF with the requested page kinds (in order)."""
    import pymupdf

    doc = pymupdf.open()
    fontfile = _font()
    for kind in pages:
        pg = doc.new_page(width=595, height=842)
        kw: dict[str, Any] = {"fontsize": 11}
        if fontfile:
            kw.update(fontfile=fontfile, fontname="synth")
        if kind in ("text", "rotated"):
            y = 72
            for _ in range(6):
                pg.insert_textbox(pymupdf.Rect(72, y, 523, y + 90), TEXT_RU + " " + TEXT_EN, **kw)
                y += 110
            if kind == "rotated":
                pg.set_rotation(90)
        elif kind == "vector":
            pg.insert_textbox(pymupdf.Rect(72, 40, 523, 120), TEXT_EN + " " + TEXT_EN, **kw)
            shape = pg.new_shape()
            for i in range(600):  # 600 separate paths (one stroke each)
                x = 72 + (i % 40) * 11
                y = 150 + (i // 40) * 40
                shape.draw_line((x, y), (x + 9, y + 30))
                shape.finish(color=(0, 0, 0), width=0.5)
            shape.commit()
        elif kind == "raster":
            pg.insert_image(pg.rect, stream=make_png(1240, 1754, text="scan"))
        elif kind == "figure":
            pg.insert_textbox(pymupdf.Rect(72, 40, 523, 200), TEXT_RU, **kw)
            pg.insert_image(pymupdf.Rect(150, 250, 450, 475), stream=make_png(600, 450))
            pg.insert_textbox(pymupdf.Rect(150, 480, 450, 520), "Рис. 1. Синтетический рисунок", **kw)
        elif kind == "mojibake":
            y = 72
            for _ in range(5):
                pg.insert_textbox(pymupdf.Rect(72, y, 523, y + 90), MOJIBAKE + " " + MOJIBAKE, fontsize=11,
                                  fontname="helv")
                y += 110
        elif kind == "broken":  # a text layer of replacement/control glyphs (Identity-H without ToUnicode look-alike)
            y = 72
            for _ in range(4):
                pg.insert_textbox(pymupdf.Rect(72, y, 523, y + 90), "�� �� " * 12, **kw)
                y += 110
        elif kind == "empty":
            pass
        else:
            raise ValueError(kind)
    if labels:
        doc.set_page_labels([{"startpage": 0, "prefix": "", "style": "r", "firstpagenum": 1}])
    data = doc.tobytes()
    doc.close()
    if bom:
        data = b"\xef\xbb\xbf" + data
    Path(path).write_bytes(data)
    return Path(path)


def make_epub(path: Path, *, n_units: int = 3) -> Path:
    """EPUB 2 with a spine, calibre page anchors, a GIF formula without alt, a JPEG figure, a table and BIBe ids."""
    from PIL import Image

    gif = io.BytesIO()
    Image.new("L", (120, 40), 255).save(gif, format="GIF")
    jpg = io.BytesIO()
    Image.new("RGB", (80, 60), (200, 100, 50)).save(jpg, format="JPEG")
    items, spine = [], []
    z = zipfile.ZipFile(path, "w")
    z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/container.xml",
               '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
               '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
               '</rootfiles></container>')
    for i in range(1, n_units + 1):
        body = [f'<h1 id="page_{2 * i - 1}">Chapter {i}</h1>', f"<p>{TEXT_EN}</p>"]
        if i == 1:
            body += ['<p class="eq"><img src="images/eq1.gif"/> (1.2)</p>',
                     '<p>Inline <img src="images/eq1.gif"/> formula inside a sentence of text.</p>',
                     '<div class="fig"><img src="images/fig1.jpg" alt="figure"/></div>',
                     "<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table>"]
        if i == n_units:
            body += ['<p id="BIBe-1">Author A. Synthetic reference one. 2020.</p>',
                     '<p id="BIBe-2">Author B. Synthetic reference two. 2021.</p>']
        body.append(f'<p><span id="page_{2 * i}"/>{TEXT_EN}</p>')
        xhtml = ('<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head>'
                 f'<title>u{i}</title></head><body>{"".join(body)}</body></html>')
        z.writestr(f"OEBPS/u{i}.xhtml", xhtml)
        items.append(f'<item id="u{i}" href="u{i}.xhtml" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="u{i}"/>')
    z.writestr("OEBPS/images/eq1.gif", gif.getvalue())
    z.writestr("OEBPS/images/fig1.jpg", jpg.getvalue())
    items += ['<item id="eq1" href="images/eq1.gif" media-type="image/gif"/>',
              '<item id="fig1" href="images/fig1.jpg" media-type="image/jpeg"/>']
    opf = ('<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Synthetic</dc:title>'
           '<dc:language>en</dc:language><dc:identifier>urn:uuid:00000000-0000-0000-0000-000000000000</dc:identifier>'
           f'</metadata><manifest>{"".join(items)}</manifest><spine>{"".join(spine)}</spine></package>')
    z.writestr("OEBPS/content.opf", opf)
    z.close()
    return Path(path)


OMML = ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>E=m</m:t></m:r><m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup>'
        '</m:sSup></m:oMath>')


def make_docx(path: Path) -> Path:
    """DOCX with a heading, paragraphs, a table and two OMML formulas (one inline, one display)."""
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    d.add_heading("Synthetic heading", level=1)
    d.add_paragraph(TEXT_RU)
    p = d.add_paragraph("Inline formula: ")
    p._p.append(parse_xml(OMML))
    disp = d.add_paragraph()
    disp._p.append(parse_xml('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                             + OMML + "</m:oMathPara>"))
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "a", "b", "1", "2"
    d.add_paragraph(TEXT_EN)
    d.save(str(path))
    return Path(path)


def make_register(resources_root: Path, rows: list[dict[str, Any]]) -> Path:
    """``00_registry/SOURCE_REGISTER.csv`` of a synthetic resources root."""
    reg = Path(resources_root) / "00_registry" / "SOURCE_REGISTER.csv"
    reg.parent.mkdir(parents=True, exist_ok=True)
    cols = ["resource_id", "canonical_path", "original_filename", "sha256", "size_bytes", "source_class",
            "evidence_scope", "priority", "scientific_role", "migration_source", "migration_status", "notes"]
    with open(reg, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})
    return reg


def register_row(source_id: str, rel: str, root: Path, **kw: Any) -> dict[str, Any]:
    p = Path(root) / rel
    row = {"resource_id": source_id, "canonical_path": rel, "original_filename": p.name,
           "sha256": sha256_file(p) if p.exists() else "0" * 64, "size_bytes": p.stat().st_size if p.exists() else 0,
           "source_class": "journal_article", "evidence_scope": "GENERAL_METHOD", "priority": "C",
           "scientific_role": "synthetic", "migration_source": "synthetic", "migration_status": "ADDED_BY_USER_EXACT",
           "notes": "synthetic"}
    row.update(kw)
    return row
